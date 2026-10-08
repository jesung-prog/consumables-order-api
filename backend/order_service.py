from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.database import OUTPUT_DIR
from backend.inventory_service import get_all_reorder_analysis


def create_order_request_excel(within_days: int = 30) -> Path | None:
    analyses = get_all_reorder_analysis()

    targets = []
    for item in analyses:
        days = item.get("days_until_reorder")
        if days is None or item.get("avg_daily_demand", 0) <= 0:
            continue
        if days <= within_days:
            order_qty = (
                item["recommended_order_qty_now"]
                if item["order_now"]
                else item["estimated_order_qty_at_trigger"]
            )
            if order_qty > 0:
                targets.append({**item, "order_qty": order_qty})

    if not targets:
        return None

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "발주요청서"

    headers = [
        "관리코드", "품목명", "단위", "현재고", "일평균수요",
        "안전일수", "추천안전재고", "재발주점", "발주까지 남은일수",
        "예상발주일", "추천발주수량", "단가", "예상금액"
    ]
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for item in targets:
        amount = item["order_qty"] * item["unit_price"]
        worksheet.append(
            [
                item["item_code"],
                item["item_name"],
                item["unit"],
                item["current_stock"],
                item["avg_daily_demand"],
                item["safety_days"],
                item["recommended_safety_stock"],
                item["reorder_point"],
                item["days_until_reorder"],
                item["predicted_order_date"],
                item["order_qty"],
                item["unit_price"],
                amount,
            ]
        )

    for row in worksheet.iter_rows(min_row=2, min_col=12, max_col=13):
        for cell in row:
            cell.number_format = '#,##0'

    for column_index, width in {
        1: 14, 2: 34, 3: 10, 4: 12, 5: 14, 6: 12, 7: 16,
        8: 12, 9: 18, 10: 14, 11: 16, 12: 14, 13: 16,
    }.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"purchase_request_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = OUTPUT_DIR / filename
    workbook.save(path)
    workbook.close()
    return path
