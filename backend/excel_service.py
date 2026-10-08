from datetime import date, datetime
from io import BytesIO
import json
import hashlib
import re
import unicodedata

from openpyxl import load_workbook


def normalize_item_name(value: str) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", "", text).lower()


def to_number(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).replace(",", "").replace("원", "").replace("₩", "").strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def parse_statement_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        raise ValueError("C5 셀에서 날짜를 찾을 수 없습니다.")

    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y년 %m월 %d일"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass

    match = re.search(r"(\d{2,4})\D+(\d{1,2})\D+(\d{1,2})", text)
    if match:
        year, month, day = map(int, match.groups())
        if year < 100:
            year += 2000
        return date(year, month, day).isoformat()

    raise ValueError(f"날짜 형식을 읽을 수 없습니다: {value}")


def read_purchase_statement_bytes(file_bytes: bytes, filename: str) -> list[dict]:
    workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    worksheet = workbook.worksheets[0]

    expected_headers = {
        "C13": "품명",
        "D13": "단위",
        "E13": "수량",
        "F13": "단가",
        "G13": "공급가액",
        "H13": "세액",
    }

    for address, expected in expected_headers.items():
        actual = worksheet[address].value
        actual_text = str(actual or "").replace(" ", "").strip()
        if expected.replace(" ", "") not in actual_text:
            raise ValueError(
                f"지원하지 않는 Excel 양식입니다. {address} 셀에 '{expected}'가 필요합니다."
            )

    statement_date = parse_statement_date(worksheet["C5"].value)
    vendor = str(worksheet["G6"].value or "").strip()

    rows = []
    for row_number in range(14, 61):
        item_name = worksheet.cell(row=row_number, column=3).value
        if item_name is None or str(item_name).strip() == "":
            continue

        quantity = to_number(worksheet.cell(row=row_number, column=5).value)
        if quantity <= 0:
            continue

        rows.append(
            {
                "source_row": row_number,
                "event_date": statement_date,
                "vendor": vendor,
                "item_name": str(item_name).strip(),
                "normalized_name": normalize_item_name(item_name),
                "unit": str(worksheet.cell(row=row_number, column=4).value or "").strip(),
                "quantity": quantity,
                "unit_price": to_number(worksheet.cell(row=row_number, column=6).value),
                "supply_amount": to_number(worksheet.cell(row=row_number, column=7).value),
                "tax_amount": to_number(worksheet.cell(row=row_number, column=8).value),
                "memo": str(worksheet.cell(row=row_number, column=9).value or "").strip(),
            }
        )

    workbook.close()

    if not rows:
        raise ValueError("Excel에서 입고 품목을 찾을 수 없습니다.")

    return rows


def create_inbound_fingerprint(rows: list[dict]) -> str:
    fingerprint_rows = [
        {
            "event_date": row["event_date"],
            "vendor": row.get("vendor", ""),
            "normalized_name": row["normalized_name"],
            "quantity": round(float(row["quantity"]), 6),
            "unit_price": round(float(row["unit_price"]), 6),
        }
        for row in rows
    ]
    fingerprint_rows.sort(
        key=lambda x: (x["event_date"], x["normalized_name"], x["quantity"], x["unit_price"])
    )
    raw = json.dumps(fingerprint_rows, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()
