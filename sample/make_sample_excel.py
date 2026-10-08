from datetime import date
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "sample_purchase_statement.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "견적(샘플)"

ws["B5"] = "날  짜 :"
ws["C5"] = date.today()
ws["F6"] = "상호"
ws["G6"] = "샘플유통"

headers = ["No.", "품명", "단위", "수량", "단가", "공급가액", "세액", "비고"]
for col, value in enumerate(headers, start=2):
    ws.cell(row=13, column=col).value = value

items = [
    (1, "샘플 고무장갑 M", "켤레", 3, 1500),
    (2, "샘플 페이퍼타올", "박스", 1, 29000),
    (3, "신규 샘플품목", "EA", 2, 5000),
]

for row_num, (no, name, unit, qty, price) in enumerate(items, start=14):
    ws.cell(row=row_num, column=2).value = no
    ws.cell(row=row_num, column=3).value = name
    ws.cell(row=row_num, column=4).value = unit
    ws.cell(row=row_num, column=5).value = qty
    ws.cell(row=row_num, column=6).value = price
    ws.cell(row=row_num, column=7).value = qty * price
    ws.cell(row=row_num, column=8).value = qty * price * 0.1

wb.save(OUTPUT)
wb.close()
print(f"Created: {OUTPUT}")
